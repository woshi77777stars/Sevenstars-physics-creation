import json

from ppt_builder import CONTENT_DIR, available_topics, build_deck, validate_deck


def test_every_deck_is_valid_json_and_has_steps():
    assert len(available_topics()) == 25
    for topic in available_topics():
        data = json.loads((CONTENT_DIR / f"{topic}.deck.json").read_text(encoding="utf-8"))
        assert data["topic"] == topic
        assert 3 <= len(data["steps"]) <= 10
        assert all(s["text"].strip() for s in data["steps"])
        assert data["title"]


def test_build_one_deck_has_native_animation_per_step(tmp_path):
    deck = build_deck("kinematics", tmp_path)
    summary = validate_deck(deck)
    content = next(s for s in summary["slides_checked"] if s.get("present"))
    steps = json.loads((CONTENT_DIR / "kinematics.deck.json").read_text(encoding="utf-8"))["steps"]
    assert content["click_effects"] == len(steps)
    assert content["body_paragraphs"] == len(steps)
    ranges = [(int(s), int(e)) for s, e in content["paragraph_ranges"]]
    assert ranges == [(i, i) for i in range(len(steps))]
    assert len(content["shape_ids"]) == 1
    assert content["issues"] == []


def test_build_every_deck_passes_animation_check(tmp_path):
    failed = []
    for topic in available_topics():
        deck = build_deck(topic, tmp_path)
        summary = validate_deck(deck)
        content = next(s for s in summary["slides_checked"] if s.get("present"))
        data = json.loads((CONTENT_DIR / f"{topic}.deck.json").read_text(encoding="utf-8"))
        if (
            content["click_effects"] != len(data["steps"])
            or content["body_paragraphs"] != len(data["steps"])
            or content.get("issues")
        ):
            failed.append(topic)
    assert failed == []


def test_deck_text_survives_round_trip(tmp_path):
    deck = build_deck("capacitance", tmp_path)
    from pptx import Presentation

    prs = Presentation(str(deck))
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.extend(p.text for p in shape.text_frame.paragraphs if p.text.strip())
    joined = "\n".join(texts)
    assert "Capacitance C = Q/V" in joined
    assert "τ = RC" in joined


def test_build_deck_uses_master_without_leaking_example_slides(tmp_path):
    deck = build_deck("waves", tmp_path)
    from pptx import Presentation

    prs = Presentation(str(deck))
    assert len(prs.slides) == 2
    texts = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                texts.extend(p.text for p in shape.text_frame.paragraphs)
    assert not any("Step one appears" in t for t in texts)
